"""
File extraction utilities for PDF, DOCX, and TXT files.
"""

import io

import fitz
import docx


def extract_text_from_pdf(file_bytes: bytes) -> str:
    """
    Extract text from a PDF file using PyMuPDF.
    """

    text_parts = []

    try:
        with fitz.open(
            stream=file_bytes,
            filetype="pdf",
        ) as pdf_doc:
            for page in pdf_doc:
                text_parts.append(page.get_text())

    except Exception as exc:
        raise ValueError(
            f"Could not read PDF file: {exc}"
        ) from exc

    return "\n".join(text_parts)


def extract_text_from_docx(file_bytes: bytes) -> str:
    """
    Extract text from DOCX paragraphs and tables.
    """

    try:
        document = docx.Document(
            io.BytesIO(file_bytes)
        )

        paragraphs = [
            paragraph.text
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        # Extract table content
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text and cell.text.strip():
                        paragraphs.append(cell.text)

        return "\n".join(paragraphs)

    except Exception as exc:
        raise ValueError(
            f"Could not read DOCX file: {exc}"
        ) from exc


def extract_text_from_txt(file_bytes: bytes) -> str:
    """
    Decode TXT files using common encodings.
    """

    for encoding in (
        "utf-8",
        "utf-8-sig",
        "latin-1",
    ):
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue

    raise ValueError(
        "Could not decode text file with common encodings."
    )


def extract_text(uploaded_file) -> str:
    """
    Extract text based on the uploaded file extension.
    """

    filename = uploaded_file.name.lower()
    file_bytes = uploaded_file.read()

    if filename.endswith(".pdf"):
        text = extract_text_from_pdf(file_bytes)

    elif filename.endswith(".docx"):
        text = extract_text_from_docx(file_bytes)

    elif filename.endswith(".txt"):
        text = extract_text_from_txt(file_bytes)

    else:
        raise ValueError(
            "Unsupported file type. Upload a PDF, DOCX, or TXT file."
        )

    if not text or not text.strip():
        raise ValueError(
            "No extractable text was found in this file."
        )

    return text